#!/usr/bin/env python3
"""Search-term database for the Goggle Rater app.

    search_terms.py seed     create search_terms.xlsx from the lists below (won't overwrite)
    search_terms.py export   write Apis_*/search_terms.txt from search_terms.xlsx

search_terms.xlsx is the place to edit. Rows with Use = N are skipped on export.
After exporting, commit and push: the app downloads the .txt files from GitHub.
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "search_terms.xlsx"
OUT = {"Health": ROOT / "Apis_health" / "search_terms.txt", "News": ROOT / "Apis_news" / "search_terms.txt"}

HEALTH = {
    "Homeopathy": [
        "homeopathy for children", "homeopathic remedies for anxiety", "classical homeopathy practitioner",
        "homeopathic first aid kit", "arnica montana uses", "homeopathy for allergies", "homeopathy research evidence",
        "constitutional homeopathy", "homeopathic remedies for colds and flu", "history of homeopathy Hahnemann",
    ],
    "Osteopathy": [
        "osteopathic manipulation benefits", "cranial osteopathy for babies", "osteopathy for back pain",
        "visceral manipulation therapy", "osteopath vs chiropractor", "osteopathy for headaches",
        "classical osteopathy Still", "osteopathy for TMJ",
    ],
    "Craniosacral": [
        "craniosacral therapy benefits", "craniosacral therapy for infants", "biodynamic craniosacral therapy",
        "craniosacral therapy for migraines", "Upledger craniosacral", "craniosacral therapy trauma",
    ],
    "Chiropractic": [
        "chiropractic adjustment benefits", "chiropractic care for children", "upper cervical chiropractic",
        "chiropractic for sciatica", "network spinal analysis", "chiropractic during pregnancy Webster technique",
        "chiropractic and immune system", "subluxation correction", "gonstead chiropractic",
    ],
    "Acupuncture & TCM": [
        "acupuncture for chronic pain", "traditional chinese medicine herbs", "acupuncture for fertility",
        "cupping therapy benefits", "moxibustion", "gua sha benefits", "TCM tongue diagnosis",
        "acupressure points for headache", "electroacupuncture", "qi gong healing", "five element acupuncture",
    ],
    "Herbal medicine": [
        "herbal medicine for beginners", "western herbalism", "medicinal herbs for immune support",
        "herbal tinctures how to make", "adaptogenic herbs", "elderberry syrup benefits", "milk thistle liver",
        "ashwagandha benefits", "berberine benefits", "turmeric curcumin research", "herbal antibiotics",
        "wild foraging medicinal plants", "Native American herbal remedies", "folk medicine remedies",
        "herbal remedies for sleep", "medicinal mushrooms lion's mane", "chaga mushroom benefits",
        "black seed oil benefits", "oregano oil uses", "neem uses",
    ],
    "Naturopathy": [
        "naturopathic doctor near me", "naturopathic medicine principles", "nature cure", "hydrotherapy healing",
        "castor oil packs", "coffee enema", "colon hydrotherapy", "liver flush", "parasite cleanse",
        "heavy metal detox", "chelation therapy", "Gerson therapy",
    ],
    "Ayurveda": [
        "ayurveda dosha diet", "ayurvedic medicine for digestion", "panchakarma detox", "oil pulling benefits",
        "triphala benefits", "ayurvedic herbs list",
    ],
    "Energy & mind-body": [
        "reiki healing", "energy medicine", "emotional freedom technique tapping", "sound healing frequencies",
        "grounding earthing benefits", "breathwork healing", "Wim Hof method", "biofield therapy",
        "PEMF therapy", "red light therapy benefits", "infrared sauna benefits", "cold plunge benefits",
        "hyperbaric oxygen therapy", "ozone therapy",
    ],
    "Bodywork": [
        "myofascial release", "lymphatic drainage massage", "rolfing structural integration", "reflexology",
        "Feldenkrais method", "Alexander technique", "trigger point therapy",
    ],
    "Ancestral nutrition": [
        "Weston Price diet", "nutrient dense foods", "organ meats benefits", "bone broth benefits",
        "fermented foods benefits", "raw milk benefits", "soaking and sprouting grains", "traditional diets health",
        "seed oils dangers", "carnivore diet results", "animal based diet", "ketogenic diet healing",
        "grass fed vs grain fed", "pastured eggs nutrition", "cod liver oil benefits", "lard and tallow cooking",
        "sugar and inflammation", "glyphosate in food", "regenerative farming food quality",
    ],
    "Functional & metabolic": [
        "functional medicine approach", "root cause medicine", "insulin resistance reversal",
        "thyroid natural treatment", "adrenal fatigue", "hormone balancing naturally",
        "bioidentical hormones", "gut microbiome healing", "leaky gut protocol", "SIBO natural treatment",
        "candida overgrowth", "mold toxicity", "lyme disease natural treatment", "methylation MTHFR",
        "chronic fatigue natural treatment", "autoimmune natural treatment", "long covid natural treatment",
        "intermittent fasting healing", "prolonged fasting autophagy",
    ],
    "Natural cancer approaches": [
        "alternative cancer treatment", "integrative oncology", "fenbendazole cancer", "ivermectin cancer",
        "high dose vitamin C cancer", "mistletoe therapy cancer", "metabolic theory of cancer",
        "Hoxsey herbal treatment", "Essiac tea",
    ],
    "Vaccines & medical freedom": [
        "vaccine injury stories", "vaccine exemption", "VAERS reports", "informed consent medicine",
        "medical freedom", "childhood vaccine schedule questions", "natural immunity research",
        "covid vaccine side effects", "ivermectin covid studies", "hydroxychloroquine studies",
    ],
    "Holistic dental & environment": [
        "holistic dentist", "biological dentistry", "fluoride dangers", "mercury amalgam removal",
        "oil pulling teeth", "EMF health effects", "water filtration fluoride", "non toxic home",
    ],
    "Women, birth & children": [
        "home birth midwife", "natural childbirth", "placenta encapsulation", "herbal remedies pregnancy",
        "breastfeeding natural remedies", "natural fertility", "natural remedies for kids fever",
        "holistic pediatrician", "ear infection natural remedy",
    ],
    "Supplements & oils": [
        "essential oils uses", "frankincense oil benefits", "vitamin D high dose", "magnesium deficiency symptoms",
        "iodine supplementation", "colloidal silver", "methylene blue benefits", "NAD+ therapy",
        "peptide therapy", "DMSO uses", "MMS chlorine dioxide",
    ],
}

NEWS = {
    "Jewish history – America": [
        "history of Jews in the United States", "Jewish immigration to America 1880-1924",
        "Jewish communities in New York history", "Jewish American history colonial era",
        "American Jewish organizations history", "Jewish influence in Hollywood history",
        "Jewish involvement in American politics", "history of Jewish lobbying in America",
        "Jews in the American civil rights movement", "history of Reform Judaism in America",
        "Jewish gangsters Prohibition era", "American Jewish philanthropy",
    ],
    "Jewish history – Europe": [
        "famous Jews in Eastern Europe", "history of Jews in Poland", "Jews in the Russian Empire Pale of Settlement",
        "history of Ashkenazi Jews", "history of Sephardic Jews", "Khazar hypothesis", "Jews in medieval Europe",
        "expulsion of Jews from Spain 1492", "Jewish Bolsheviks history", "Jews in the Soviet Union",
        "Rothschild family history", "history of Jews in Germany", "Jewish population before and after World War II",
        "history of antisemitism", "Holocaust history", "Jewish refugees after WWII",
    ],
    "Ukraine": [
        "Jewish influence in Ukraine", "Jewish community in Ukraine history", "Zelensky Jewish heritage",
        "Kolomoisky Ukraine oligarch", "Ukraine oligarchs", "Ukraine war news", "Ukraine peace talks",
        "Chabad in Ukraine", "Uman Rosh Hashanah pilgrimage",
    ],
    "Torah, Talmud & Judaism": [
        "rabbinical Torah studies", "what does the Talmud say about non-Jews", "Talmud teachings on gentiles",
        "Noahide laws", "Babylonian Talmud explained", "Kabbalah history", "Chabad Lubavitch beliefs",
        "Orthodox Judaism beliefs", "Jewish messianism", "Jewish law halacha", "Mishnah and Gemara",
        "Third Temple movement", "red heifer Temple Mount", "Zohar teachings", "Judaism and Christianity differences",
        "Pharisees and rabbinic Judaism", "Jewish views on Jesus",
    ],
    "Israel – policy & politics": [
        "Israeli political policy", "Israeli government coalition news", "Netanyahu news", "Israeli settlements West Bank",
        "Greater Israel plans", "Yinon plan", "Israeli judicial reform", "Israel nation state law",
        "Israel military aid from US", "Israeli intelligence Mossad operations", "Israel nuclear weapons",
        "Israeli politics religious parties", "Ben Gvir Smotrich policies", "Israel Palestine two state solution",
        "how was the State of Israel established", "Balfour Declaration history", "Nakba 1948 history",
        "Zionism history", "Christian Zionism", "Israel lobby AIPAC", "Israel and US Congress",
    ],
    "Middle East news": [
        "Middle Eastern news", "War in Lebanon news", "Hezbollah news", "Gaza news", "Gaza ceasefire talks",
        "West Bank news", "US Iran negotiations", "Iran nuclear program news", "Syria news", "Yemen Houthis Red Sea",
        "Iraq news", "Saudi Arabia Israel normalization", "Abraham Accords", "Qatar mediation news",
        "Turkey Middle East policy", "Egypt Gaza border", "Jordan news", "Gulf states news",
    ],
    "Immigration & activism": [
        "Jewish immigration activism", "HIAS refugee resettlement", "Jewish organizations immigration policy",
        "Jewish groups and civil rights law", "ADL activism history", "Jewish voices for peace", "anti-Zionist Jews",
        "Neturei Karta",
    ],
    "Religion & power": [
        "Peter Thiel antichrist", "Peter Thiel antichrist lectures", "Peter Thiel politics", "Palantir surveillance",
        "dispensationalism end times Israel", "evangelicals and Israel", "Catholic Church and Judaism relations",
        "Vatican Israel relations", "Freemasonry history", "Bohemian Grove",
    ],
    "Media & geopolitics": [
        "independent journalism Middle East", "media bias Israel coverage", "censorship of Israel criticism",
        "free speech antisemitism laws", "foreign policy realism", "neoconservatives Iraq war history",
        "Epstein Maxwell intelligence ties", "USS Liberty incident", "Lavon affair", "Israel spying on US",
    ],
}


def seed():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    if XLSX.exists() and "--force" not in sys.argv:
        sys.exit(f"{XLSX.name} already exists; it's yours to edit now. Use --force to overwrite.")
    wb = Workbook()
    info = wb.active
    info.title = "Read me"
    for row in [
        ["Goggle Rater search terms"],
        [],
        ["Each time the app needs new sites, it picks a random term from the matching sheet and searches Brave."],
        ["Edit freely: add rows, change categories, set Use to N to switch a term off without deleting it."],
        ["Then run:  python tools/search_terms.py export   and push. The app picks the terms up on its next start."],
    ]:
        info.append(row)
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 110

    honey = PatternFill("solid", fgColor="FCEA7E")
    for name, data in (("Health", HEALTH), ("News", NEWS)):
        ws = wb.create_sheet(name)
        ws.append(["Term", "Category", "Use", "Notes"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = honey
        seen = set()
        for cat, terms in data.items():
            for t in terms:
                if t.lower() not in seen:
                    seen.add(t.lower())
                    ws.append([t, cat, "Y", ""])
        for col, w in zip("ABCD", (48, 28, 7, 40)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:D{ws.max_row}"
        dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(f"C2:C{ws.max_row + 500}")
        for row in ws.iter_rows(min_row=2, max_col=3, min_col=3):
            row[0].alignment = Alignment(horizontal="center")
        print(f"{name}: {ws.max_row - 1} terms")
    wb.save(XLSX)
    print(f"wrote {XLSX}")


def export():
    from openpyxl import load_workbook

    wb = load_workbook(XLSX, read_only=True)
    for name, out in OUT.items():
        terms, seen = [], set()
        for row in wb[name].iter_rows(min_row=2, values_only=True):
            term = (row[0] or "").strip() if row else ""
            use = (str(row[2]).strip().upper() if len(row) > 2 and row[2] is not None else "Y")
            if term and use != "N" and term.lower() not in seen:
                seen.add(term.lower())
                terms.append(term)
        out.write_text("\n".join(terms) + "\n", encoding="utf-8", newline="\n")
        print(f"{out.relative_to(ROOT)}: {len(terms)} terms")


if __name__ == "__main__":
    {"seed": seed, "export": export}.get(sys.argv[1] if len(sys.argv) > 1 else "", lambda: sys.exit(__doc__))()
